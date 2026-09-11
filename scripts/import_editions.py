"""Import edition-pinned public workbooks. Run with --download to retrieve them again."""
import argparse, hashlib, json, math, re, urllib.request
from datetime import datetime, timezone
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
SOURCES=ROOT/'data'/'sources'
URLS={'happiness2026.xlsx':'https://files.worldhappiness.report/WHR26_Data_Figure_2.1.xlsx','sdg2026.xlsx':'https://dashboards.sdgindex.org/static/downloads/database_2026.xlsx','dst2026.pdf':'https://dst.gov.in/sites/default/files/ST%20INDICATORS%20TABLES%202025-26.pdf'}
IDS={'India':'IND','Singapore':'SGP','Japan':'JPN','China':'CHN','United States':'USA','Denmark':'DNK','Norway':'NOR','Sweden':'SWE','Finland':'FIN','Iceland':'ISL'}
GOALS=['No poverty','Zero hunger','Good health and wellbeing','Quality education','Gender equality','Clean water and sanitation','Affordable and clean energy','Decent work and economic growth','Industry, innovation and infrastructure','Reduced inequalities','Sustainable cities and communities','Responsible consumption and production','Climate action','Life below water','Life on land','Peace, justice and strong institutions','Partnerships for the goals']
def finite(v):return isinstance(v,(float,int)) and not isinstance(v,bool) and math.isfinite(v)
def records(sheet):
 it=sheet.iter_rows(values_only=True); headers=next(it)
 return [dict(zip(headers,r)) for r in it]
def main(download=False):
 SOURCES.mkdir(parents=True,exist_ok=True)
 if download:
  for name,url in URLS.items():
   with urllib.request.urlopen(url,timeout=90) as r:(SOURCES/name).write_bytes(r.read())
 provenance={name:{'url':url,'retrievedAt':datetime.fromtimestamp((SOURCES/name).stat().st_mtime,timezone.utc).isoformat(),'sha256':hashlib.sha256((SOURCES/name).read_bytes()).hexdigest()} for name,url in URLS.items()}
 happiness=[]
 w=openpyxl.load_workbook(SOURCES/'happiness2026.xlsx',read_only=True,data_only=True)
 for r in records(w.worksheets[0]):
  if r.get('Country name') not in IDS or not finite(r.get('Life evaluation (3-year average)')):continue
  y=int(r['Year']);happiness.append({'country':IDS[r['Country name']],'year':y,'value':r['Life evaluation (3-year average)'],'low':r['Lower whisker'],'high':r['Upper whisker'],'footnote':f'World Happiness Report 2026; three-year survey average ending {y}, not a single-year result. '+(f'95% confidence interval: {r["Lower whisker"]} to {r["Upper whisker"]}.' if finite(r['Lower whisker']) and finite(r['Upper whisker']) else 'Confidence interval not supplied in this workbook for this observation.')})
 w.close()
 w=openpyxl.load_workbook(SOURCES/'sdg2026.xlsx',read_only=True,data_only=True)
 back=[r for r in records(w['Backdated SDG Index']) if r.get('id') in IDS.values()]
 current={r['Country Code ISO3']:r for r in records(w['SDR2026 Data']) if r.get('Country Code ISO3') in IDS.values()}
 codebook=records(w['Codebook'])
 sdg=[]
 for code,r in current.items():
  goals=[]
  for n,name in enumerate(GOALS,1):
   history=[{'year':int(o['year']),'value':o[f'goal{n}']} for o in back if o['id']==code and finite(o.get(f'goal{n}'))]
   details=[]
   for m in codebook:
    if m.get('SDG')!=n:continue
    key=m.get('IndCode'); val=r.get(m.get('Indicator'))
    details.append({'id':key,'name':m.get('Indicator'),'value':val if finite(val) else None,'year':r.get('Year: '+key),'imputed':r.get('Imputation: '+key),'status':r.get('Dashboard Color: '+key),'trend':r.get('Trend: '+key),'source':m.get('Source'),'url':m.get('Dwldlink'),'definition':m.get('Description')})
   goals.append({'number':n,'name':name,'status':r.get(f'Goal {n} Dash'),'trend':r.get(f'Goal {n} Trend'),'history':history,'indicators':details})
  sdg.append({'country':code,'reportYear':2026,'score':r['2026 SDG Index Score'],'rank':r['2026 SDG Index Rank'],'goals':goals,'history':[{'year':int(o['year']),'value':o['sdgi_s']} for o in back if o['id']==code and finite(o.get('sdgi_s'))]})
 w.close()
 series=[{'id':'WHR_LIFE','name':'Happiness / life evaluation','officialName':'Life evaluation (3-year average), World Happiness Report 2026','category':'Wellbeing','source':'WHR','unit':'Cantril ladder, 0–10','direction':'higher','url':'https://www.worldhappiness.report/data-sharing/','apiUrl':URLS['happiness2026.xlsx'],'organisation':'University of Oxford Wellbeing Research Centre / Gallup','definition':'Average self-reported life evaluation on the Cantril ladder. Each data point is a three-year average labelled by its ending survey year. Confidence intervals are preserved in the footnotes. This is not an index constructed by adding GDP or other explanatory factors.','sourceUpdated':'World Happiness Report 2026 edition','retrievedAt':provenance['happiness2026.xlsx']['retrievedAt'],'status':'ok','observations':happiness},
 {'id':'SDSN_INDEX','name':'SDG Index (backdated, 2026 edition)','officialName':'SDSN backdated SDG Index, SDR 2026 edition','category':'SDG progress','source':'SDSN','unit':'Score, 0–100','direction':'higher','url':'https://dashboards.sdgindex.org/downloads/','apiUrl':URLS['sdg2026.xlsx'],'organisation':'Sustainable Development Solutions Network (SDSN)','definition':'Historical SDG Index recalculated in the 2026 report edition. A score of 100 corresponds to the report methodology targets, not proof that all UN SDG targets have been met. SDSN is complementary to, and distinct from, the official UN indicator framework.','sourceUpdated':'Sustainable Development Report 2026 edition','retrievedAt':provenance['sdg2026.xlsx']['retrievedAt'],'status':'ok','observations':[{'country':o['id'],'year':int(o['year']),'value':o['sdgi_s'],'footnote':'Backdated series from one edition (SDR 2026). Historical estimates use the report methodology and may include imputations.'} for o in back if finite(o.get('sdgi_s'))]}]
 # These report tables were transcribed and checked against the source PDF.
 research={'country':'IND','edition':'DST S&T Indicators Tables 2025–26','url':URLS['dst2026.pdf'],'retrievedAt':provenance['dst2026.pdf']['retrievedAt'],
 'disciplines':{'reference':'1 April 2024','table':23,'pdfPage':48,'unit':'Full-time personnel (headcount)','exclusion':'Survey responses across sectors excluding higher education. This is personnel specialisation, not a count of laboratories or a comprehensive national infrastructure inventory.','total':268623,'rows':[{'name':n,'male':m,'female':f,'value':t} for n,m,f,t in [('Agricultural sciences',20886,6383,27269),('Engineering & technology',115702,25442,141144),('Medical sciences',44464,17446,61910),('Natural sciences',23248,9280,32528),('Social sciences',3201,2571,5772)]]},
 'sectors':{'reference':'2023–24 financial year','table':2,'pdfPage':19,'unit':'₹ crore, current prices','total':244767.81,'rows':[{'name':n,'value':v} for n,v in [('Central government institutions',75655.72),('State government institutions',9796.10),('Public higher education',22353.61),('Private higher education',8554.92),('Scientific & industrial research organisations',7798.28),('Private industry',110553.40),('Public industry',10055.78)]]},
 'states':{'reference':'2023–24 financial year','table':9,'pdfPage':27,'unit':'₹ crore, current prices','total':9796.10,'exclusion':'State-government R&D including state agricultural universities. Excludes central and private spending located in each state. Omitted states/UTs are not assigned zero by this dashboard.','rows':[{'name':n,'value':v} for n,v in [('Andhra Pradesh',154.09),('Assam',169.02),('Bihar',152.73),('Chhattisgarh',354.59),('Gujarat',814.79),('Haryana',595.03),('Himachal Pradesh',374.68),('Jammu & Kashmir',153.03),('Jharkhand',94.09),('Karnataka',1130.15),('Kerala',326.46),('Madhya Pradesh',155.90),('Maharashtra',528.37),('Manipur',.38),('Meghalaya',9.17),('Odisha',320.60),('Punjab',994.78),('Rajasthan',640.65),('Tamil Nadu',1214.50),('Telangana',652.02),('Uttar Pradesh',433.99),('Uttarakhand',220.34),('West Bengal',306.74)]]}}
 out={'importedAt':datetime.now(timezone.utc).isoformat(),'provenance':provenance,'series':series,'sdg':sdg,'research':research}
 (ROOT/'data'/'editions.json').write_text(json.dumps(out,ensure_ascii=False,allow_nan=False),encoding='utf-8')
 print('Imported',len(happiness),'happiness observations;',len(sdg),'SDG country profiles; DST tables 2, 9 and 23.')
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--download',action='store_true');main(p.parse_args().download)

