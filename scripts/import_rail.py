from pathlib import Path
from openpyxl import load_workbook
import json,hashlib
p=Path('data/states.json');d=json.loads(p.read_text());lookup={s['name'].upper():s['name'] for s in d['states']};lookup.update({'JAMMU AND KASHMIR':'Jammu & Kashmir','UTTARAKHAND':'Uttarakhand'})
f=next(Path('data/sources').glob('4_*.xlsx'));sheet=load_workbook(f,data_only=True).worksheets[0];rows=list(sheet.values);headers=rows[3];dest=None;flows=[]
for row in rows[6:]:
 if row[1] in lookup:dest=lookup[row[1]]
 if not dest or not row[2] or 'TOTAL' in str(row[2]).upper():continue
 for j,head in enumerate(headers):
  if head not in lookup:continue
  origin=lookup[head];value=row[j]
  if origin!=dest and isinstance(value,(int,float)):
   assert value>=0
   flows.append({'from':origin,'to':dest,'commodity':row[2],'value':value})
# Independently reconcile the matrix with the official state-commodity totals (Table II).
totals=load_workbook(next(Path('data/sources').glob('2_*.xlsx')),data_only=True).worksheets[0];r=list(totals.values);heads=r[3];checks=0
for row in r[6:]:
 if not isinstance(row[1],(int,float)) or not row[2]:continue
 for j,head in enumerate(heads):
  if head not in lookup:continue
  for offset,key in [(0,'from'),(1,'to')]:
   value=row[j+offset]
   if isinstance(value,(int,float)):
    actual=sum(f['value'] for f in flows if f[key]==lookup[head] and f['commodity']==row[2]);assert abs(value-actual)<.1,(head,row[2],key,value,actual);checks+=1
assert flows and checks>100
source='https://www.dgciskol.gov.in/pub_inland.aspx';d['rail']={'period':'2024-25','unit':'Tonnes','source':source,'table':'Table IV, Principal Commodities by Rail','commodities':sorted(set(f['commodity'] for f in flows)),'flows':flows,'reconciledCells':checks}
d['provenance']=[q for q in d['provenance'] if not q['file'].startswith('4_')]+[{'file':f.name,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}]
p.write_text(json.dumps(d,ensure_ascii=True),encoding='utf-8');print('Rail',len(flows),'flows;',len(d['rail']['commodities']),'commodities;',checks,'independent reconciliation checks')
